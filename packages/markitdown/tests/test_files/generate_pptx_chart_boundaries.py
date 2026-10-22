"""Rebuild the chart/text, chart/table, and chart-only regression slides.

Run with python-pptx installed; the tests use the saved PPTX directly.
"""

from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def main() -> None:
    presentation = Presentation()
    for following in ("text", "table", None):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Revenue", (10, 20))
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            0,
            0,
            Inches(5),
            Inches(3),
            data,
        )
        if following == "text":
            slide.shapes.add_textbox(
                0, Inches(4), Inches(5), Inches(1)
            ).text = "Follow-up analysis."
        elif following == "table":
            table = slide.shapes.add_table(
                2, 2, 0, Inches(4), Inches(5), Inches(1)
            ).table
            for cell, value in zip(
                (
                    table.cell(0, 0),
                    table.cell(0, 1),
                    table.cell(1, 0),
                    table.cell(1, 1),
                ),
                ("Item", "Details", "A", "B"),
            ):
                cell.text = value

    presentation.save(Path(__file__).with_name("pptx_chart_boundaries.pptx"))


if __name__ == "__main__":
    main()
