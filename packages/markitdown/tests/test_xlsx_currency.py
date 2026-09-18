"""Currency-formatted Excel cells keep their label (microsoft/markitdown#53)."""

import io

import openpyxl
import pytest

from markitdown import StreamInfo
from markitdown.converters import XlsxConverter


_INFO = StreamInfo(extension=".xlsx")


def _workbook() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["Item", "Count", "Cost", "Weight", "Total"])
    sheet.append(["Breakfast", 20, 5, None, 100])
    sheet.append(["Laptops", 5, 1199, None, 5995])
    sheet.append(["Car tires", 8, 199, "150 kg", 1592])
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = '"$"#,##0.00'
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _convert(data: bytes) -> str:
    return XlsxConverter().convert(io.BytesIO(data), _INFO).markdown


def test_currency_cells_keep_their_label() -> None:
    markdown = _convert(_workbook())
    assert "$5" in markdown
    assert "$1199" in markdown
    assert "$100" in markdown
    assert "Breakfast" in markdown
    assert "150 kg" in markdown


def test_plain_cells_are_untouched() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Count"])
    sheet.append(["Breakfast", 20])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    markdown = _convert(stream.getvalue())
    assert "| Breakfast | 20 |" in markdown
    assert "$" not in markdown


def test_euro_suffix_format() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Price"])
    sheet.append([42])
    sheet["A2"].number_format = '#,##0.00 [$€-x-euro2]'
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    assert "42€" in _convert(stream.getvalue())
