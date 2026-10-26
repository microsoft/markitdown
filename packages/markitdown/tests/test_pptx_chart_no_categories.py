"""PowerPoint charts that have no category axis (XY scatter and bubble).

Those charts keep their data points inside the series as ``(x, y)`` pairs, so
``chart.plots[0].categories`` is empty and the converter used to emit a table
that ended right after its separator line, dropping every value.
"""

import io

import pptx
from pptx.chart.data import BubbleChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from markitdown import MarkItDown, StreamInfo

SCATTER_VALUES = [10.0, 20.0, 30.0]


def _chart_presentation(chart_type, chart_data) -> io.BytesIO:
    """Build a one-slide deck holding a single chart titled "Sales"."""
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    chart = slide.shapes.add_chart(
        chart_type, Inches(1), Inches(1), Inches(6), Inches(4), chart_data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Sales"
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _scatter_presentation() -> io.BytesIO:
    """Build a one-slide deck holding an XY (scatter) chart."""
    chart_data = XyChartData()
    series = chart_data.add_series("Revenue")
    for idx, value in enumerate(SCATTER_VALUES, start=1):
        series.add_data_point(float(idx), value)
    return _chart_presentation(XL_CHART_TYPE.XY_SCATTER, chart_data)


def _bubble_presentation() -> io.BytesIO:
    """Build a one-slide deck holding a bubble chart."""
    chart_data = BubbleChartData()
    series = chart_data.add_series("Revenue")
    for idx, value in enumerate(SCATTER_VALUES, start=1):
        series.add_data_point(float(idx), value, 1.0)
    return _chart_presentation(XL_CHART_TYPE.BUBBLE, chart_data)


def _convert(stream: io.BytesIO) -> str:
    return (
        MarkItDown()
        .convert_stream(stream, stream_info=StreamInfo(extension=".pptx"))
        .markdown
    )


def test_pptx_scatter_chart_keeps_data_points() -> None:
    markdown = _convert(_scatter_presentation())

    assert "### Chart: Sales" in markdown
    for idx, value in enumerate(SCATTER_VALUES, start=1):
        assert f"| {idx} | {value} |" in markdown


def test_pptx_bubble_chart_keeps_data_points() -> None:
    markdown = _convert(_bubble_presentation())

    assert "### Chart: Sales" in markdown
    for idx, value in enumerate(SCATTER_VALUES, start=1):
        assert f"| {idx} | {value} |" in markdown
