"""Regression tests for literal delimiters in PowerPoint chart tables."""

import io

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from markitdown import StreamInfo
from markitdown.converters import PptxConverter


def _build_chart_stream(categories, series):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series:
        chart_data.add_series(name, values)
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(8),
        Inches(5),
        chart_data,
    )
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return stream


def _convert_chart(categories, series):
    stream = _build_chart_stream(categories, series)
    return PptxConverter().convert(stream, StreamInfo(extension=".pptx")).markdown


@pytest.mark.parametrize(
    ("category", "expected_category"),
    [
        ("literal | pipe", "literal \\| pipe"),
        *[
            (
                f"{count} backslashes {chr(92) * count}| pipe",
                f"{count} backslashes {chr(92) * (2 * count + 1)}| pipe",
            )
            for count in (1, 2, 3)
        ],
        (f"first line{chr(10)}second line", "first line<br>second line"),
        (
            f"first line{chr(13)}{chr(10)}second line",
            "first line<br>second line",
        ),
        (
            f"first line{chr(13)}second line",
            "first line<br>second line",
        ),
    ],
)
def test_chart_table_escapes_literal_delimiters(category, expected_category):
    markdown = _convert_chart(
        [category],
        [("Series | name", (0,)), ("Plain", (-2,))],
    )

    assert markdown.splitlines() == [
        "<!-- Slide number: 1 -->",
        "",
        "",
        "### Chart",
        "",
        "| Category | Series \\| name | Plain |",
        "|---|---|---|",
        f"| {expected_category} | 0.0 | -2.0 |",
    ]


def test_chart_table_preserves_plain_numeric_output():
    markdown = _convert_chart(
        ["Plain"],
        [("Plain", (0,)), ("Other", (-2,))],
    )

    assert markdown.splitlines() == [
        "<!-- Slide number: 1 -->",
        "",
        "",
        "### Chart",
        "",
        "| Category | Plain | Other |",
        "|---|---|---|",
        "| Plain | 0.0 | -2.0 |",
    ]
