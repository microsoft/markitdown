#!/usr/bin/env python3 -m pytest
"""Reading-order regression tests for the PPTX converter.

Shapes are read top-to-bottom and left-to-right within a row. A strict
``sorted(key=(top, left))`` sorts by ``top`` first, so a shape that shares a row
with a neighbor but starts a hair higher (the right column of a two-column body,
a right-aligned value next to its label) is emitted before its left-hand
neighbor. These tests pin the left-to-right behavior.
"""
from typing import Any

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches

from markitdown import MarkItDown
from markitdown.converters._pptx_converter import _sort_shapes_reading_order


def _textbox(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
) -> Any:
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    box.text_frame.text = text
    return box


def _convert(tmp_path: Any, build: Any) -> str:
    prs = Presentation()
    build(prs, prs.slide_layouts[6])  # layout 6 is blank
    deck = tmp_path / "reading_order.pptx"
    prs.save(str(deck))
    return MarkItDown().convert(str(deck)).markdown


def test_two_column_body_reads_left_then_right(tmp_path: Any) -> None:
    """The right column's box starts slightly higher, but must still be read second."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        _textbox(slide, 0.5, 1.5, 4, 3, "LEFT column body text")
        _textbox(slide, 5.0, 1.4, 4, 3, "RIGHT column body text")  # 0.1" higher

    md = _convert(tmp_path, build)
    assert md.index("LEFT column body text") < md.index("RIGHT column body text")


def test_same_row_label_reads_before_right_aligned_value(tmp_path: Any) -> None:
    """A right-aligned value that sits a hair above its label is still read after it."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        _textbox(slide, 0.5, 1.5, 4, 0.5, "Job Title")
        _textbox(slide, 6.0, 1.45, 3, 0.5, "2020 - Present")  # slightly higher

    md = _convert(tmp_path, build)
    assert md.index("Job Title") < md.index("2020 - Present")


def test_stacked_single_column_order_is_preserved(tmp_path: Any) -> None:
    """Non-overlapping, stacked shapes keep their natural top-to-bottom order."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        _textbox(slide, 0.5, 3.0, 8, 0.5, "third line")
        _textbox(slide, 0.5, 1.0, 8, 0.5, "first line")
        _textbox(slide, 0.5, 2.0, 8, 0.5, "second line")

    md = _convert(tmp_path, build)
    assert md.index("first line") < md.index("second line") < md.index("third line")


def _add_grid(slide: Any, origin_left: float) -> None:
    for r in range(3):
        for c in range(3):
            _textbox(
                slide, origin_left + c * 2.3, 0.6 + r * 2.1, 2.0, 1.5, f"cell {r}{c}"
            )


def test_grid_beside_left_sidebar_reads_row_major(tmp_path: Any) -> None:
    """A full-height sidebar must not pull the grid beside it into column-major order."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        _textbox(slide, 0.3, 0.5, 1.6, 6.5, "SIDEBAR")
        _add_grid(slide, 2.5)

    md = _convert(tmp_path, build)
    assert md.index("SIDEBAR") < md.index("cell 00")
    assert md.index("cell 00") < md.index("cell 01") < md.index("cell 02")
    assert md.index("cell 02") < md.index("cell 10")  # row 0 fully before row 1


def test_grid_beside_right_sidebar_reads_before_the_sidebar(tmp_path: Any) -> None:
    """A right-hand full-height sidebar is read after the grid, which stays row-major."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        _add_grid(slide, 0.5)
        _textbox(slide, 7.6, 0.5, 1.6, 6.5, "SIDEBAR")

    md = _convert(tmp_path, build)
    assert md.index("cell 00") < md.index("cell 01") < md.index("cell 02")
    assert md.index("cell 02") < md.index("cell 10")
    assert md.index("cell 22") < md.index("SIDEBAR")  # sidebar read last


def test_full_width_title_overlapping_the_row_below_reads_first(tmp_path: Any) -> None:
    """A full-width title whose box overlaps the top row is read before that row,
    even when a column starts a little to its left."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        # The title spans the slide and dips into the row below, so it bands with
        # the two columns. The left column starts a hair further left, so a plain
        # left-to-right sort within that band would emit it before the title.
        _textbox(slide, 0.5, 0.3, 9.0, 1.6, "FULL WIDTH TITLE")
        _textbox(slide, 0.4, 1.0, 4.0, 3.0, "LEFT column")
        _textbox(slide, 5.0, 1.0, 4.0, 3.0, "RIGHT column")

    md = _convert(tmp_path, build)
    assert md.index("FULL WIDTH TITLE") < md.index("LEFT column")
    assert md.index("LEFT column") < md.index("RIGHT column")


def test_full_width_caption_under_a_tall_figure_reads_after_it(tmp_path: Any) -> None:
    """A full-width caption nested in the lower half of a taller, narrower figure
    is read after the figure, not before it because its box starts further left."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        _textbox(slide, 1.0, 2.0, 3.0, 3.3, "FIGURE body content")
        _textbox(slide, 0.3, 4.3, 9.0, 0.6, "Caption spanning the width")

    md = _convert(tmp_path, build)
    assert md.index("FIGURE body content") < md.index("Caption spanning the width")


def test_diagonal_staircase_does_not_chain_into_one_row(tmp_path: Any) -> None:
    """A diagonal staircase must not chain into a single band. Adjacent steps
    overlap ~60% (so they share a row), but a step and the one two below it
    overlap only ~20% (below the same-row threshold). Banding on the growing
    union would stretch the band down the whole staircase and read it as one
    scrambled line; the top pair of steps must stay ahead of the bottom pair."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        # Each step is one row down and to the left of the previous one.
        _textbox(slide, 6.0, 0.5, 3.0, 1.0, "step A top")
        _textbox(slide, 4.0, 0.9, 3.0, 1.0, "step B")
        _textbox(slide, 2.0, 1.3, 3.0, 1.0, "step C")
        _textbox(slide, 0.0, 1.7, 3.0, 1.0, "step D bottom")

    md = _convert(tmp_path, build)
    top_pair = max(md.index("step A top"), md.index("step B"))
    bottom_pair = min(md.index("step C"), md.index("step D bottom"))
    assert top_pair < bottom_pair


class _StubShape:
    """Minimal stand-in exposing the geometry attributes the helper reads."""

    def __init__(
        self,
        top: Any,
        left: Any,
        width: Any = 100,
        height: Any = 100,
        name: str = "",
    ) -> None:
        self.top = top
        self.left = left
        self.width = width
        self.height = height
        self.name = name


def test_none_top_sorts_first_and_zero_top_is_a_real_position() -> None:
    """top=None sorts first (like upstream's -inf); top=0 is a real coordinate, not missing."""
    none_top = _StubShape(top=None, left=100, name="none_top")
    zero_top = _StubShape(top=0, left=100, name="zero_top")
    lower = _StubShape(top=200, left=100, name="lower")
    order = [s.name for s in _sort_shapes_reading_order([lower, zero_top, none_top])]
    assert order == ["none_top", "zero_top", "lower"]


def test_none_left_sorts_first_within_a_row() -> None:
    """left=None sorts before a positioned shape in the same row (upstream -inf)."""
    none_left = _StubShape(top=50, left=None, name="none_left")
    positioned = _StubShape(top=50, left=100, name="positioned")
    order = [s.name for s in _sort_shapes_reading_order([positioned, none_left])]
    assert order == ["none_left", "positioned"]


def test_none_height_shape_is_kept_not_dropped() -> None:
    """A shape with height=None (treated as zero) is still emitted, never dropped."""
    normal = _StubShape(top=10, left=0, height=100, name="normal")
    no_height = _StubShape(top=300, left=0, height=None, name="no_height")
    order = [s.name for s in _sort_shapes_reading_order([no_height, normal])]
    assert order == ["normal", "no_height"]


def test_group_shape_children_are_read_in_reading_order(tmp_path: Any) -> None:
    """Shapes inside a group get the same reading-order rule (the converter
    recurses into group children), so a right child that starts slightly higher
    is still read after its left neighbor."""

    def build(prs: Any, blank: Any) -> None:
        slide = prs.slides.add_slide(blank)
        group = slide.shapes.add_group_shape()
        left = group.shapes.add_textbox(
            Inches(0.5), Inches(1.5), Inches(3), Inches(0.5)
        )
        left.text_frame.text = "GROUP left cell"
        right = group.shapes.add_textbox(
            Inches(5.0), Inches(1.4), Inches(3), Inches(0.5)
        )
        right.text_frame.text = "GROUP right cell"

    md = _convert(tmp_path, build)
    assert md.index("GROUP left cell") < md.index("GROUP right cell")


if __name__ == "__main__":
    import sys

    sys.exit(pytest.main([__file__, "-v"]))
