import io

from openpyxl import Workbook

from markitdown import MarkItDown, StreamInfo


def _xlsx_bytes(workbook: Workbook) -> bytes:
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


def test_completely_empty_sheet_is_skipped() -> None:
    """A sheet with no columns must not emit malformed table syntax."""
    workbook = Workbook()
    workbook.active.title = "HasData"
    workbook.active["A1"] = "col"
    workbook.active["A2"] = "v"
    workbook.create_sheet("CompletelyEmpty")

    result = MarkItDown().convert_stream(
        io.BytesIO(_xlsx_bytes(workbook)),
        stream_info=StreamInfo(extension=".xlsx"),
    )

    assert "## HasData" in result.markdown
    assert "| col |" in result.markdown
    assert "| v |" in result.markdown
    assert "CompletelyEmpty" not in result.markdown
    assert "|\n|  |" not in result.markdown


def test_header_only_sheet_keeps_empty_table() -> None:
    """A sheet with a header row but no data already produces a valid table."""
    workbook = Workbook()
    workbook.active.title = "HeaderOnly"
    workbook.active["A1"] = "col"

    result = MarkItDown().convert_stream(
        io.BytesIO(_xlsx_bytes(workbook)),
        stream_info=StreamInfo(extension=".xlsx"),
    )

    assert "## HeaderOnly" in result.markdown
    assert "| col |" in result.markdown
    assert "| --- |" in result.markdown


def test_single_completely_empty_workbook_emits_nothing() -> None:
    workbook = Workbook()
    workbook.active.title = "Empty"

    result = MarkItDown().convert_stream(
        io.BytesIO(_xlsx_bytes(workbook)),
        stream_info=StreamInfo(extension=".xlsx"),
    )

    assert result.markdown.strip() == ""
    assert "|" not in result.markdown
