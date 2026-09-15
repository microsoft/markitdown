"""Chart categories must retain the context that distinguishes their values."""

import io

import pytest

from markitdown import MarkItDown, StreamInfo

pptx = pytest.importorskip("pptx")
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def _convert_chart(data: CategoryChartData) -> str:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(8),
        Inches(5),
        data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Sales"
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return (
        MarkItDown()
        .convert_stream(stream, stream_info=StreamInfo(extension=".pptx"))
        .markdown
    )


def test_chart_retains_parent_categories_and_value_order() -> None:
    data = CategoryChartData()
    first_year = data.add_category("2024")
    first_year.add_sub_category("Q1")
    first_year.add_sub_category("Q2")
    data.add_category("2025").add_sub_category("Q1")
    data.add_series("Actual", (10, 20, 30))
    data.add_series("Budget", (11, 22, 33))

    markdown = _convert_chart(data)

    assert "### Chart: Sales" in markdown
    assert markdown.splitlines()[-5:] == [
        "| Category | Actual | Budget |",
        "|---|---|---|",
        "| 2024 / Q1 | 10.0 | 11.0 |",
        "| 2024 / Q2 | 20.0 | 22.0 |",
        "| 2025 / Q1 | 30.0 | 33.0 |",
    ]


def test_chart_retains_all_category_levels() -> None:
    data = CategoryChartData()
    for country in ("US", "Canada"):
        region = data.add_category(country).add_sub_category("West")
        region.add_sub_category("Retail")
        region.add_sub_category("Online")
    data.add_series("Actual", (1, 2, 3, 4))

    assert _convert_chart(data).splitlines()[-4:] == [
        "| US / West / Retail | 1.0 |",
        "| US / West / Online | 2.0 |",
        "| Canada / West / Retail | 3.0 |",
        "| Canada / West / Online | 4.0 |",
    ]


@pytest.mark.parametrize("categories", [("Q1", "Q2"), (2024, 2025)])
def test_flat_chart_categories_keep_existing_output(categories) -> None:
    data = CategoryChartData()
    data.categories = categories
    data.add_series("Actual", (10, 20))

    assert _convert_chart(data).splitlines()[-2:] == [
        f"| {categories[0]} | 10.0 |",
        f"| {categories[1]} | 20.0 |",
    ]
