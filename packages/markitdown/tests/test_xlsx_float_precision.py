"""Spreadsheet float values must survive conversion without pandas' 6-digit scientific notation."""
import io

from openpyxl import Workbook

from markitdown import MarkItDown, StreamInfo


def test_xlsx_float_precision_preserved():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "money"
    ws["A2"] = 123456789.123
    ws["A3"] = 0.1
    ws["A4"] = 1e-10
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    result = MarkItDown().convert_stream(buf, stream_info=StreamInfo(extension=".xlsx"))
    assert "123456789.123" in result.markdown
    assert "1.234568e+08" not in result.markdown
    assert "| 0.1 |" in result.markdown
