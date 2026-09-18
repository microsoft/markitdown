"""Regression test: charts with missing data points must not render "None".

A series value can be None when the underlying worksheet cell is blank, and a
series can also be shorter than the category list. Both cases used to end up
in the Markdown table as the literal string "None".
"""

import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from markitdown import MarkItDown, StreamInfo


def _presentation_with_sparse_chart() -> io.BytesIO:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])

    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2", "Q3"]
    chart_data.add_series("Series 1", (1.0, None, 3.0))  # blank data point
    chart_data.add_series("Series 2", (0.5,))  # shorter than the categories
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def test_chart_missing_values_are_empty_cells() -> None:
    result = MarkItDown().convert(
        _presentation_with_sparse_chart(),
        stream_info=StreamInfo(
            mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            extension=".pptx",
        ),
    )
    table_lines = [
        line for line in result.markdown.splitlines() if line.startswith("|")
    ]
    assert len(table_lines) == 5  # header, separator, three data rows
    # Missing points become empty cells, not the string "None".
    assert table_lines[2] == "| Q1 | 1.0 | 0.5 |"
    assert table_lines[3] == "| Q2 |  |  |"
    assert table_lines[4] == "| Q3 | 3.0 |  |"
    assert "None" not in result.markdown
