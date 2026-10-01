#!/usr/bin/env python3 -m pytest
"""Robustness test for the PPTX converter.

A shape whose type python-pptx cannot identify raises NotImplementedError when
its ``shape_type`` is read. One such shape must not abort the whole
presentation, so the type checks treat it as "none of the handled types" and
fall through instead of propagating.

The fixture is a real ``<p:sp>`` with no placeholder, no preset or custom
geometry, and a ``<p:cNvSpPr>`` without ``txBox="1"`` - exactly the element
python-pptx cannot map - injected into a slide's shape tree.
"""
import io
from typing import Any

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches

from markitdown import MarkItDown
from markitdown._stream_info import StreamInfo
from markitdown.converters._pptx_converter import PptxConverter

# Not a placeholder, no <a:prstGeom>/<a:custGeom>, and cNvSpPr lacks txBox="1",
# so Shape.shape_type raises NotImplementedError. It still carries a text frame,
# so falling through must recover its text rather than drop it.
_UNRECOGNIZED_SP = (
    '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    '<p:nvSpPr><p:cNvPr id="9999" name="BadShape"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
    "<p:spPr/>"
    "<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r>"
    "<a:t>text in an unrecognized shape</a:t>"
    "</a:r></a:p></p:txBody>"
    "</p:sp>"
)


def _slide_with_unrecognized_shape() -> tuple[Any, Any]:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "surrounding text"
    slide.shapes._spTree.append(parse_xml(_UNRECOGNIZED_SP))
    return prs, slide


def _bad_shape(slide: Any) -> Any:
    return next(s for s in slide.shapes if s.name == "BadShape")


def test_fixture_really_is_unrecognized() -> None:
    """Guard the fixture itself: python-pptx must raise on its shape_type."""
    _, slide = _slide_with_unrecognized_shape()
    with pytest.raises(NotImplementedError):
        _ = _bad_shape(slide).shape_type


def test_shape_type_or_none_swallows_unrecognized_type() -> None:
    _, slide = _slide_with_unrecognized_shape()
    assert PptxConverter._shape_type_or_none(_bad_shape(slide)) is None


def test_unrecognized_shape_is_neither_picture_nor_table() -> None:
    _, slide = _slide_with_unrecognized_shape()
    converter = PptxConverter()
    bad = _bad_shape(slide)
    assert converter._is_picture(bad) is False
    assert converter._is_table(bad) is False


def test_one_unrecognized_shape_does_not_abort_the_deck() -> None:
    """The deck still converts, and both the surrounding text and the unrecognized
    shape's own text are recovered (falling through reaches the text-frame branch)."""
    prs, _ = _slide_with_unrecognized_shape()
    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    result = MarkItDown().convert_stream(
        buffer, stream_info=StreamInfo(extension=".pptx")
    )
    assert "surrounding text" in result.markdown
    assert "text in an unrecognized shape" in result.markdown


if __name__ == "__main__":
    import sys

    sys.exit(pytest.main([__file__, "-v"]))
