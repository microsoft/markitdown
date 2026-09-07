"""
Unit tests for PptxConverterWithOCR.

For each PPTX test file: convert with a mock OCR service then compare the
full output string against the expected snapshot.

OCR block format used by the converter:
    *[Image OCR]
    MOCK_OCR_TEXT_12345
    [End OCR]*

Note: PPTX slide text uses literal backslash-n (\\n) sequences from the
underlying PPTX converter template; OCR blocks use real newlines.
"""

import io
import sys
from pathlib import Path
from typing import Any

import pptx
import pptx.shapes.autoshape
import pptx.slide
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from markitdown_ocr._ocr_service import OCRResult  # noqa: E402
from markitdown_ocr._pptx_converter_with_ocr import (  # noqa: E402
    PptxConverterWithOCR,
)
from markitdown import StreamInfo  # noqa: E402

TEST_DATA_DIR = Path(__file__).parent / "ocr_test_data"

_MOCK_TEXT = "MOCK_OCR_TEXT_12345"
_OCR_BLOCK = f"*[Image OCR]\n{_MOCK_TEXT}\n[End OCR]*"


class MockOCRService:
    def extract_text(
        self,  # noqa: ANN101
        image_stream: Any,
        **kwargs: Any,
    ) -> OCRResult:
        return OCRResult(text=_MOCK_TEXT, backend_used="mock")


@pytest.fixture(scope="module")
def svc() -> MockOCRService:
    return MockOCRService()


def _convert(filename: str, ocr_service: MockOCRService) -> str:
    path = TEST_DATA_DIR / filename
    if not path.exists():
        pytest.skip(f"Test file not found: {path}")
    converter = PptxConverterWithOCR()
    with open(path, "rb") as f:
        return converter.convert(
            f, StreamInfo(extension=".pptx"), ocr_service=ocr_service
        ).text_content


# ---------------------------------------------------------------------------
# pptx_image_start.pptx
# ---------------------------------------------------------------------------


def test_pptx_image_start(svc: MockOCRService) -> None:
    # Slide 1: title "Welcome" followed by an image
    expected = (
        "\\n\\n<!-- Slide number: 1 -->\\n# Welcome\\n\\n"
        "\n*[Image OCR]\nMOCK_OCR_TEXT_12345\n[End OCR]*"
    )
    assert _convert("pptx_image_start.pptx", svc) == expected


# ---------------------------------------------------------------------------
# pptx_image_middle.pptx
# ---------------------------------------------------------------------------


def test_pptx_image_middle(svc: MockOCRService) -> None:
    # Slide 1: Introduction | Slide 2: Architecture + image | Slide 3: Conclusion  # noqa: E501
    expected = (
        "\\n\\n<!-- Slide number: 1 -->\\n# Introduction"
        "\\n\\n\\n\\n<!-- Slide number: 2 -->\\n# Architecture\\n\\n"
        "\n*[Image OCR]\nMOCK_OCR_TEXT_12345\n[End OCR]*"
        "\\n\\n<!-- Slide number: 3 -->\\n# Conclusion\\n\\n"
    )
    assert _convert("pptx_image_middle.pptx", svc) == expected


# ---------------------------------------------------------------------------
# pptx_image_end.pptx
# ---------------------------------------------------------------------------


def test_pptx_image_end(svc: MockOCRService) -> None:
    # Slide 1: Presentation | Slide 2: Thank You + image
    expected = (
        "\\n\\n<!-- Slide number: 1 -->\\n# Presentation"
        "\\n\\n\\n\\n<!-- Slide number: 2 -->\\n# Thank You\\n\\n"
        "\n*[Image OCR]\nMOCK_OCR_TEXT_12345\n[End OCR]*"
    )
    assert _convert("pptx_image_end.pptx", svc) == expected


# ---------------------------------------------------------------------------
# pptx_multiple_images.pptx
# ---------------------------------------------------------------------------


def test_pptx_multiple_images(svc: MockOCRService) -> None:
    # Slide 1: two images, no title text
    expected = (
        "\\n\\n<!-- Slide number: 1 -->\\n# \\n"
        "\n*[Image OCR]\nMOCK_OCR_TEXT_12345\n[End OCR]*"
        "\n\n*[Image OCR]\nMOCK_OCR_TEXT_12345\n[End OCR]*"
    )
    assert _convert("pptx_multiple_images.pptx", svc) == expected


# ---------------------------------------------------------------------------
# pptx_complex_layout.pptx
# ---------------------------------------------------------------------------


def test_pptx_complex_layout(svc: MockOCRService) -> None:
    expected = (
        "\\n\\n<!-- Slide number: 1 -->\\n# Product Comparison"
        "\\n\\nOur products lead the market\\n"
        "\n*[Image OCR]\nMOCK_OCR_TEXT_12345\n[End OCR]*"
    )
    assert _convert("pptx_complex_layout.pptx", svc) == expected


# ---------------------------------------------------------------------------
# No OCR service — no OCR tags emitted
# ---------------------------------------------------------------------------


def test_pptx_no_ocr_service_no_tags() -> None:
    path = TEST_DATA_DIR / "pptx_image_middle.pptx"
    if not path.exists():
        pytest.skip(f"Test file not found: {path}")
    converter = PptxConverterWithOCR()
    with open(path, "rb") as f:
        md = converter.convert(f, StreamInfo(extension=".pptx")).text_content
    assert "*[Image OCR]" not in md
    assert "[End OCR]*" not in md


def test_pptx_ocr_chart_no_title_text_frame() -> None:
    from unittest.mock import MagicMock
    from markitdown_ocr._pptx_converter_with_ocr import PptxConverterWithOCR

    # python-pptx's ChartTitle.text_frame is destructive -- it creates a text
    # frame if one isn't already present, so it never returns None.
    # has_text_frame is the property that actually reflects presence/absence.
    mock_chart = MagicMock()
    mock_chart.has_title = True
    mock_chart.chart_title.has_text_frame = False

    mock_category = MagicMock()
    mock_category.label = "Cat 1"
    mock_chart.plots = [MagicMock(categories=[mock_category])]

    mock_series = MagicMock()
    mock_series.name = "Series 1"
    mock_series.values = [10.0]
    mock_chart.series = [mock_series]

    converter = PptxConverterWithOCR()
    result = converter._convert_chart_to_markdown(mock_chart)

    assert "### Chart" in result
    assert "Cat 1" in result
    assert "Series 1" in result
    assert ":" not in result


def test_pptx_ocr_chart_with_title_text_frame() -> None:
    from unittest.mock import MagicMock
    from markitdown_ocr._pptx_converter_with_ocr import PptxConverterWithOCR

    mock_chart = MagicMock()
    mock_chart.has_title = True
    mock_chart.chart_title.has_text_frame = True
    mock_chart.chart_title.text_frame.text = "Revenue"

    mock_category = MagicMock()
    mock_category.label = "Cat 1"
    mock_chart.plots = [MagicMock(categories=[mock_category])]

    mock_series = MagicMock()
    mock_series.name = "Series 1"
    mock_series.values = [10.0]
    mock_chart.series = [mock_series]

    converter = PptxConverterWithOCR()
    result = converter._convert_chart_to_markdown(mock_chart)

    assert "### Chart: Revenue" in result


def test_pptx_ocr_none_shape_text(
    svc: MockOCRService, monkeypatch: pytest.MonkeyPatch
) -> None:
    # python-pptx returns None for shape.text when an <a:r> run has no <a:t>
    # child, which happens in some third-party-generated decks.
    monkeypatch.setattr(
        pptx.shapes.autoshape.Shape,
        "text",
        property(lambda self: None),
        raising=True,
    )

    assert _convert("pptx_complex_layout.pptx", svc) is not None


def test_pptx_ocr_none_notes_text(monkeypatch: pytest.MonkeyPatch) -> None:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "some notes"
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    real_notes = pptx.slide.NotesSlide.notes_text_frame

    class _NoneText:
        text = None

    def fake_notes(self: pptx.slide.NotesSlide) -> Any:
        frame = real_notes.fget(self)
        if frame is None:
            return None
        return _NoneText()

    monkeypatch.setattr(
        pptx.slide.NotesSlide,
        "notes_text_frame",
        property(fake_notes),
        raising=True,
    )

    converter = PptxConverterWithOCR()
    result = converter.convert(buf, StreamInfo(extension=".pptx"))

    assert result.text_content is not None
